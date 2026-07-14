# -*- coding: utf-8 -*-
"""
按「同义词导入模板.xlsx」生成 0721ACGN 同义词导入表：
  - A：标准词（每行一条，不出现空单元格）
  - B：同义词（英文分号 ; 分隔，**至多 5 条**，多则截断，不拆多行）
默认每次保存到 **新文件名**（带时间戳），不覆盖旧的 xlsx。
数据来源：0721acgn/docs/专业词库-0721ACGN.md（第八节一般不导入）。
"""
from __future__ import annotations

import re
from collections import OrderedDict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

DOCS = Path(__file__).resolve().parent
MD_PATH = DOCS / "专业词库-0721ACGN.md"
OUTPUT_PREFIX = "0721ACGN-synonym-import"

# 默认每次运行写入带时间戳的新文件，避免覆盖已打开的或未备份的旧表。

# 每个标准词一行；B 列同义词总数不超过该值，超出部分丢弃（不拆行、A 列不留空）。
SYNONYMS_MAX_PER_CELL = 5

TEMPLATE_CANDIDATES = [
    Path(r"c:\Users\MoQi\Downloads\同义词导入模板 .xlsx"),
    Path(r"c:\Users\MoQi\Downloads\同义词导入模板.xlsx"),
]


def split_zh_list(s: str) -> list[str]:
    s = re.sub(r"[（(][^）)]*[）)]", " ", s)
    parts = re.split(r"[、，,]+", s)
    return [p.strip() for p in parts if p.strip()]


def norm_std(key: str) -> tuple[str, list[str]]:
    key = key.strip()
    if " / " in key:
        parts = [p.strip() for p in key.split(" / ") if p.strip()]
        return parts[0], parts[1:]
    return key, []


def extract_markdown_tables(section_body: str) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in section_body.splitlines():
        line_st = line.strip()
        if not line_st.startswith("|"):
            continue
        if re.match(r"^\|\s*-+", line_st):
            continue
        cells = [c.strip() for c in line_st.strip("|").split("|")]
        rows.append(cells)
    return rows


def get_section(md: str, heading_prefix: str) -> str:
    pat = "^" + re.escape(heading_prefix) + r"[^\n]*\n(.*?)(?=^## |\Z)"
    m = re.search(pat, md, re.MULTILINE | re.DOTALL)
    return m.group(1) if m else ""


def extract_fenced_block(section_body: str) -> str:
    m = re.search(r"```\s*\r?\n(.*?)```", section_body, re.DOTALL)
    return (m.group(1) if m else "").strip()


def collect_terms(md_text: str) -> OrderedDict[str, list[str]]:
    od: OrderedDict[str, list[str]] = OrderedDict()

    def add(std: str, syns: list[str]) -> None:
        std = std.strip()
        if not std:
            return
        if std not in od:
            od[std] = []
        bucket = od[std]
        for x in syns:
            t = x.strip()
            if not t or t == std:
                continue
            if t not in bucket:
                bucket.append(t)

    sec1 = get_section(md_text, "## 一、")
    for row in extract_markdown_tables(sec1):
        if len(row) < 2:
            continue
        if row[0] == "标准词":
            continue
        std0, tail = norm_std(row[0])
        add(std0, split_zh_list(row[1]) + tail)

    sec2 = get_section(md_text, "## 二、")
    for row in extract_markdown_tables(sec2):
        if len(row) < 2:
            continue
        if row[0] == "标准能力 / 控件":
            continue
        std0, tail = norm_std(row[0])
        add(std0, split_zh_list(row[1]) + tail)

    sec3 = get_section(md_text, "## 三、")
    for row in extract_markdown_tables(sec3):
        if len(row) < 2:
            continue
        if row[0] == "标准概念":
            continue
        std0, tail = norm_std(row[0])
        add(std0, split_zh_list(row[1]) + tail)

    # 四：文化俗称
    sec4 = get_section(md_text, "## 四、")
    rows4 = extract_markdown_tables(sec4)
    skip_hdr = {"用户说法"}
    for row in rows4:
        if len(row) < 2:
            continue
        if row[0] in skip_hdr:
            continue
        left = row[0].strip()

        parts = split_zh_list(left)
        if "、" in left and len(parts) >= 2:
            add(parts[0], parts[1:])
            add(parts[1], parts[:1])
            continue
        if "/" in left and ("相册" in left or "返图" in left):
            std0, tail = norm_std(left.replace(" ", ""))
            add(std0, tail)
            continue
        if "巴士" in left or "tour" in left.lower():
            std0, tail = norm_std(re.sub(r"\s+", " ", left))
            add(std0, tail)
            continue
        if left == "打卡":
            add("打卡", ["打点", "到站拍照", "纪念"])
            continue
        if left == "取景":
            add("取景", [])
            continue
        std0, tail = norm_std(left)
        add(std0, tail)

    # 五：地名口述
    sec5 = get_section(md_text, "## 五、")
    for row in extract_markdown_tables(sec5):
        if len(row) < 2:
            continue
        if row[0] == "用户说法":
            continue
        lhs = row[0]
        if any(x in lhs for x in ("首都圈", "关东", "关西")):
            add("日本大区", split_zh_list(lhs))
            continue
        if "电车" in lhs or "JR" in lhs:
            add("日系铁路", split_zh_list(lhs))
            continue
        add("日本", split_zh_list(lhs))

    # 六：代码块
    sec6 = get_section(md_text, "## 六、")
    for line in extract_fenced_block(sec6).splitlines():
        line = line.strip()
        if ":" not in line:
            continue
        k, v = line.split(":", 1)
        k = k.strip()
        syns = [piece.strip() for piece in re.split(r"[,，]", v) if piece.strip()]
        add(k, syns)

    # 七：归一
    sec7 = get_section(md_text, "## 七、")
    for raw in extract_fenced_block(sec7).splitlines():
        line = raw.strip()
        if "→" not in line:
            continue
        left, right = line.split("→", 1)
        users = [p.strip() for p in re.split(r"\s*/\s*", left) if p.strip()]
        right_clean = (
            right.strip().split("（", 1)[0].strip().split("(", 1)[0].strip()
        )
        if any(x in right_clean for x in ("忽略", "轻量互动")):
            continue
        stds = [p.strip() for p in re.split(r"\s*/\s*", right_clean) if p.strip()]
        for std in stds:
            add(std, users)

    return od


def cleanup(od: OrderedDict[str, list[str]]) -> None:
    for std in list(od.keys()):
        seen: set[str] = set()
        out: list[str] = []
        for s in od[std]:
            if s != std and s not in seen:
                seen.add(s)
                out.append(s)
        od[std] = out


def load_template_workbook():
    for p in TEMPLATE_CANDIDATES:
        if p.is_file():
            return load_workbook(p), p
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "标准词"
    ws["B1"] = "同义词"
    ws["B2"] = (
        f"英文分号(;)分隔；每个标准词一行；B 列同义词至多{SYNONYMS_MAX_PER_CELL}个，超出请改词库后重导。"
    )
    return wb, None


def main() -> Path:
    terms = collect_terms(MD_PATH.read_text(encoding="utf-8"))
    cleanup(terms)

    wb, tpl_path = load_template_workbook()
    ws = wb.active

    if tpl_path:
        addon = (
            f"每标准词单行；同义词至多{SYNONYMS_MAX_PER_CELL}个（;分隔）；勿留空标准词。"
        )
        b2_prev = ws["B2"].value
        prev = "" if b2_prev is None else str(b2_prev).strip()
        ws["B2"] = addon if not prev else f"{prev} {addon}"

    while ws.max_row > 2:
        ws.delete_rows(3)

    limit = SYNONYMS_MAX_PER_CELL
    rows_written = 0
    truncated_terms = 0
    for std, syns in sorted(terms.items(), key=lambda x: x[0]):
        if not syns:
            continue
        if len(syns) > limit:
            truncated_terms += 1
        clipped = syns[:limit]
        ws.append([std, ";".join(clipped)])
        rows_written += 1

    out_path = DOCS / (
        f"{OUTPUT_PREFIX}-{datetime.now():%Y%m%d-%H%M%S}.xlsx"
    )
    wb.save(out_path)
    std_with_syns = sum(1 for sy in terms.values() if sy)
    trunc_note = (
        f"；有 {truncated_terms} 条标准词的同义词被截断至{limit}个"
        if truncated_terms
        else ""
    )
    print(("模版: " + str(tpl_path)) if tpl_path else "模版: （内置表头）")
    print(
        f"已写入新文件: {out_path}  （数据行: {rows_written}，有同义词的标准词: {std_with_syns}{trunc_note}）"
    )
    return out_path


if __name__ == "__main__":
    main()
